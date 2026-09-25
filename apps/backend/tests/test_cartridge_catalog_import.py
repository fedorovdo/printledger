import io
import os
import re
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from fastapi.routing import APIRoute
from openpyxl import Workbook
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.orm import Session

from app.api.auth import require_admin
from app.api.cartridge_catalog import router as cartridge_catalog_router
from app.api.catalog import patch_cartridge_model, post_cartridge_model
from app.db.base import Base
from app.models import (
    CartridgeInventoryTransaction,
    CartridgeModel,
    PrinterInstalledCartridge,
)
from app.models.enums import CartridgeType
from app.schemas.catalog import CartridgeModelCreate, CartridgeModelUpdate
from app.services.cartridge_catalog_excel import (
    MAX_CATALOG_NOTES_LENGTH,
    CatalogApplyValidationError,
    CatalogExcelError,
    CatalogSnapshotConflict,
    apply_cartridge_catalog_import,
    preview_cartridge_catalog_import,
)
from app.services.cartridge_inventory_excel import (
    ACTUAL_NEW_HEADER,
    ACTUAL_REFILLED_HEADER,
    COMMENT_HEADER,
    INVENTORY_SHEET_NAME,
    MIN_STOCK_HEADER,
    MODEL_NAME_HEADER,
    SKU_HEADER,
    TYPE_HEADER,
    VENDOR_HEADER,
    preview_inventory_import,
)


CATALOG_HEADERS = (
    VENDOR_HEADER,
    MODEL_NAME_HEADER,
    SKU_HEADER,
    TYPE_HEADER,
    MIN_STOCK_HEADER,
    COMMENT_HEADER,
)


def _workbook_bytes(
    rows: list[tuple[object, ...]],
    headers: tuple[str, ...] = CATALOG_HEADERS,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = INVENTORY_SHEET_NAME
    worksheet["A1"] = "Служебная строка"
    worksheet.append([])
    worksheet.append([])
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _without_dimension(content: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(content), "r")
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data, substitutions = re.subn(
                    rb"<dimension\s+ref=\"[^\"]+\"\s*/>",
                    b"",
                    data,
                    count=1,
                )
                if substitutions != 1:
                    raise AssertionError("worksheet dimension was not found")
            target.writestr(item, data)
    return output.getvalue()


class CartridgeCatalogImportTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite+pysqlite:///:memory:")
        Base.metadata.create_all(
            self.engine,
            tables=[
                CartridgeModel.__table__,
                CartridgeInventoryTransaction.__table__,
                PrinterInstalledCartridge.__table__,
            ],
        )
        self.db = Session(self.engine)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def _add_model(
        self,
        model_name: str,
        *,
        vendor: str | None = "Vendor",
        purchase_sku: str | None = None,
        cartridge_type: CartridgeType = CartridgeType.toner,
        min_stock_level: int = 3,
        notes: str | None = "original",
    ) -> CartridgeModel:
        model = CartridgeModel(
            vendor=vendor,
            model_name=model_name,
            purchase_sku=purchase_sku,
            cartridge_type=cartridge_type,
            min_stock_level=min_stock_level,
            notes=notes,
            is_active=True,
        )
        self.db.add(model)
        self.db.commit()
        self.db.refresh(model)
        return model

    def _preview(self, rows: list[tuple[object, ...]]):
        return preview_cartridge_catalog_import(self.db, _workbook_bytes(rows))

    def test_unknown_model_is_create(self):
        preview = self._preview([("HP", "CF259A", "SKU-1", "Тонер", 4, "note")])

        self.assertEqual(preview.summary.create_rows, 1)
        self.assertEqual(preview.rows[0].status, "create")
        self.assertIsNone(preview.rows[0].existing_model_id)

    def test_existing_exact_model_is_existing(self):
        model = self._add_model("CF259A", vendor="HP", purchase_sku="SKU-1")

        preview = self._preview([(" hp ", " cf259a ", "sku-1", "toner", 99, "new")])

        self.assertEqual(preview.rows[0].status, "existing")
        self.assertEqual(preview.rows[0].existing_model_id, model.id)

    def test_duplicate_normalized_names_block_all_rows(self):
        preview = self._preview(
            [
                ("A", "CE505X", "SKU-A", "toner", 0, None),
                ("B", " ce505x ", "SKU-B", "toner", 0, None),
            ]
        )

        self.assertEqual(preview.summary.error_rows, 2)
        self.assertTrue(all("DUPLICATE MODEL" in row.errors for row in preview.rows))

    def test_sku_collision_with_existing_model_is_error(self):
        self._add_model("OLD", purchase_sku="SHARED")

        preview = self._preview([("A", "NEW", " shared ", "toner", 0, None)])

        self.assertEqual(preview.rows[0].status, "error")
        self.assertTrue(any("SKU уже используется" in error for error in preview.rows[0].errors))

    def test_duplicate_sku_for_new_rows_blocks_both(self):
        preview = self._preview(
            [
                ("A", "MODEL-A", "SHARED", "toner", 0, None),
                ("B", "MODEL-B", " shared ", "ink", 0, None),
            ]
        )

        self.assertTrue(all(row.status == "error" for row in preview.rows))
        self.assertTrue(all("DUPLICATE SKU" in row.errors for row in preview.rows))

    def test_blank_sku_is_not_a_duplicate(self):
        preview = self._preview(
            [
                ("A", "MODEL-A", None, "toner", 0, None),
                ("B", "MODEL-B", "  ", "ink", 0, None),
            ]
        )

        self.assertTrue(all(row.status == "create" for row in preview.rows))
        self.assertTrue(all("DUPLICATE SKU" not in row.errors for row in preview.rows))

    def test_supported_type_mapping(self):
        for excel_value, expected in (
            ("Тонер", CartridgeType.toner),
            (" toner ", CartridgeType.toner),
            ("Чернила", CartridgeType.ink),
            ("INK", CartridgeType.ink),
            ("Другое", CartridgeType.other),
            ("other", CartridgeType.other),
        ):
            with self.subTest(excel_value=excel_value):
                preview = self._preview(
                    [("Vendor", f"MODEL-{expected.value}-{excel_value}", None, excel_value, 0, None)]
                )
                self.assertEqual(preview.rows[0].cartridge_type, expected)
                self.assertEqual(preview.rows[0].status, "create")

    def test_unknown_type_is_error(self):
        preview = self._preview([("A", "MODEL", None, "powder", 0, None)])

        self.assertEqual(preview.rows[0].status, "error")
        self.assertTrue(any(TYPE_HEADER in error for error in preview.rows[0].errors))

    def test_blank_min_stock_defaults_to_zero(self):
        preview = self._preview([("A", "MODEL", None, "toner", None, None)])

        self.assertEqual(preview.rows[0].min_stock_level, 0)
        self.assertEqual(preview.rows[0].status, "create")

    def test_explicit_zero_min_stock_is_zero(self):
        preview = self._preview([("A", "MODEL", None, "toner", 0, None)])

        self.assertEqual(preview.rows[0].min_stock_level, 0)
        self.assertEqual(preview.rows[0].status, "create")

    def test_invalid_min_stock_values_are_errors(self):
        for value in (-1, 1.5, "one", True):
            with self.subTest(value=value):
                preview = self._preview([("A", f"MODEL-{value}", None, "toner", value, None)])
                self.assertEqual(preview.rows[0].status, "error")
                self.assertTrue(any(MIN_STOCK_HEADER in error for error in preview.rows[0].errors))

    def test_existing_model_conflicting_identifiers_are_errors(self):
        self._add_model("MODEL", vendor="Vendor A", purchase_sku="SKU-A")

        vendor_conflict = self._preview(
            [("Vendor B", "MODEL", "SKU-A", "toner", 0, None)]
        )
        sku_conflict = self._preview(
            [("Vendor A", "MODEL", "SKU-B", "toner", 0, None)]
        )

        self.assertEqual(vendor_conflict.rows[0].status, "error")
        self.assertEqual(sku_conflict.rows[0].status, "error")

    def test_existing_model_conflicting_type_is_error(self):
        self._add_model("MODEL", cartridge_type=CartridgeType.toner)

        preview = self._preview(
            [("Vendor", "MODEL", None, "Другое", 0, None)]
        )

        self.assertEqual(preview.rows[0].status, "error")
        self.assertTrue(
            any("Тип не совпадает" in error for error in preview.rows[0].errors)
        )

    def test_existing_model_with_blank_excel_identifiers_is_not_overwritten(self):
        model = self._add_model("MODEL", vendor="Vendor", purchase_sku="SKU")

        preview = self._preview([(None, "MODEL", None, "toner", 99, "replacement")])
        response = apply_cartridge_catalog_import(
            self.db,
            _workbook_bytes([(None, "MODEL", None, "toner", 99, "replacement")]),
            preview.summary.snapshot_hash,
        )
        self.db.refresh(model)

        self.assertEqual(response.models_created, 0)
        self.assertEqual(model.vendor, "Vendor")
        self.assertEqual(model.purchase_sku, "SKU")
        self.assertEqual(model.cartridge_type, CartridgeType.toner)
        self.assertEqual(model.min_stock_level, 3)
        self.assertEqual(model.notes, "original")

    def test_existing_empty_db_identifier_only_warns_and_is_not_updated(self):
        model = self._add_model("MODEL", vendor=None, purchase_sku=None)

        preview = self._preview([("Excel Vendor", "MODEL", "EXCEL-SKU", "toner", 0, None)])
        self.assertEqual(preview.rows[0].status, "existing")
        self.assertEqual(len(preview.rows[0].warnings), 2)

        apply_cartridge_catalog_import(
            self.db,
            _workbook_bytes([("Excel Vendor", "MODEL", "EXCEL-SKU", "toner", 0, None)]),
            preview.summary.snapshot_hash,
        )
        self.db.refresh(model)
        self.assertIsNone(model.vendor)
        self.assertIsNone(model.purchase_sku)

    def test_validation_error_blocks_whole_apply(self):
        content = _workbook_bytes(
            [
                ("A", "VALID", "SKU-1", "toner", 0, None),
                ("B", "INVALID", "SKU-2", "unknown", 0, None),
            ]
        )
        preview = preview_cartridge_catalog_import(self.db, content)

        with self.assertRaises(CatalogApplyValidationError):
            apply_cartridge_catalog_import(self.db, content, preview.summary.snapshot_hash)

        self.assertEqual(self.db.scalar(select(func.count(CartridgeModel.id))), 0)

    def test_mid_batch_failure_rolls_back_all_models(self):
        content = _workbook_bytes(
            [
                ("A", "FIRST", "SKU-1", "toner", 0, None),
                ("B", "FAIL", "SKU-2", "ink", 0, None),
            ]
        )
        preview = preview_cartridge_catalog_import(self.db, content)

        def fail_second(_mapper, _connection, target):
            if target.model_name == "FAIL":
                raise RuntimeError("simulated insert failure")

        event.listen(CartridgeModel, "before_insert", fail_second)
        try:
            with self.assertRaisesRegex(RuntimeError, "simulated insert failure"):
                apply_cartridge_catalog_import(
                    self.db,
                    content,
                    preview.summary.snapshot_hash,
                )
        finally:
            event.remove(CartridgeModel, "before_insert", fail_second)

        self.assertEqual(self.db.scalar(select(func.count(CartridgeModel.id))), 0)

    def test_mid_batch_failure_is_rolled_back_for_fresh_session(self):
        self.db.close()
        self.engine.dispose()
        with tempfile.TemporaryDirectory() as directory:
            database_path = os.path.join(directory, "catalog-test.sqlite")
            engine = create_engine(f"sqlite+pysqlite:///{database_path}")
            Base.metadata.create_all(engine, tables=[CartridgeModel.__table__])
            db = Session(engine)
            content = _workbook_bytes(
                [
                    ("A", "FIRST", "SKU-1", "toner", 0, None),
                    ("B", "SECOND", "SKU-2", "ink", 0, None),
                    ("C", "FAIL", "SKU-3", "other", 0, None),
                ]
            )
            preview = preview_cartridge_catalog_import(db, content)

            def fail_third(_mapper, _connection, target):
                if target.model_name == "FAIL":
                    raise RuntimeError("simulated third insert failure")

            event.listen(CartridgeModel, "before_insert", fail_third)
            try:
                with self.assertRaisesRegex(
                    RuntimeError,
                    "simulated third insert failure",
                ):
                    apply_cartridge_catalog_import(
                        db,
                        content,
                        preview.summary.snapshot_hash,
                    )
            finally:
                event.remove(CartridgeModel, "before_insert", fail_third)
                db.close()

            with Session(engine) as fresh_db:
                count = fresh_db.scalar(select(func.count(CartridgeModel.id)))
            engine.dispose()

        self.assertEqual(count, 0)

    def test_catalog_change_makes_snapshot_stale(self):
        content = _workbook_bytes([("A", "NEW", "SKU-1", "toner", 0, None)])
        preview = preview_cartridge_catalog_import(self.db, content)
        self._add_model("OTHER", purchase_sku="OTHER-SKU")

        with self.assertRaisesRegex(CatalogSnapshotConflict, "Каталог изменился"):
            apply_cartridge_catalog_import(
                self.db,
                content,
                preview.summary.snapshot_hash,
            )

        self.assertIsNone(
            self.db.scalar(select(CartridgeModel).where(CartridgeModel.model_name == "NEW"))
        )

    def test_second_apply_with_old_hash_is_stale(self):
        content = _workbook_bytes([("A", "NEW", "SKU-1", "toner", 0, None)])
        preview = preview_cartridge_catalog_import(self.db, content)

        first = apply_cartridge_catalog_import(
            self.db,
            content,
            preview.summary.snapshot_hash,
        )
        with self.assertRaises(CatalogSnapshotConflict):
            apply_cartridge_catalog_import(
                self.db,
                content,
                preview.summary.snapshot_hash,
            )

        self.assertEqual(first.models_created, 1)
        self.assertEqual(self.db.scalar(select(func.count(CartridgeModel.id))), 1)

    def test_apply_creates_no_stock_transactions(self):
        content = _workbook_bytes([("A", "NEW", "SKU-1", "toner", 5, None)])
        preview = preview_cartridge_catalog_import(self.db, content)
        before = self.db.scalar(select(func.count(CartridgeInventoryTransaction.id)))

        apply_cartridge_catalog_import(self.db, content, preview.summary.snapshot_hash)
        after = self.db.scalar(select(func.count(CartridgeInventoryTransaction.id)))

        self.assertEqual((before, after), (0, 0))

    def test_snapshot_is_deterministic_and_changes_with_file_or_catalog(self):
        content = _workbook_bytes([("A", "NEW", "SKU-1", "toner", 5, "note")])
        first = preview_cartridge_catalog_import(self.db, content)
        second = preview_cartridge_catalog_import(self.db, content)
        changed_file = preview_cartridge_catalog_import(
            self.db,
            _workbook_bytes([("A", "NEW", "SKU-1", "toner", 6, "note")]),
        )
        self._add_model("OTHER", purchase_sku="OTHER-SKU")
        changed_catalog = preview_cartridge_catalog_import(self.db, content)

        self.assertEqual(first.summary.snapshot_hash, second.summary.snapshot_hash)
        self.assertNotEqual(first.summary.snapshot_hash, changed_file.summary.snapshot_hash)
        self.assertNotEqual(first.summary.snapshot_hash, changed_catalog.summary.snapshot_hash)

    def test_snapshot_changes_for_each_imported_field(self):
        base_row = ("Vendor", "MODEL", "SKU", "toner", 5, "note")
        base_hash = self._preview([base_row]).summary.snapshot_hash
        variants = (
            ("Other", "MODEL", "SKU", "toner", 5, "note"),
            ("Vendor", "MODEL-2", "SKU", "toner", 5, "note"),
            ("Vendor", "MODEL", "SKU-2", "toner", 5, "note"),
            ("Vendor", "MODEL", "SKU", "ink", 5, "note"),
            ("Vendor", "MODEL", "SKU", "toner", 6, "note"),
            ("Vendor", "MODEL", "SKU", "toner", 5, "changed"),
        )

        for row in variants:
            with self.subTest(row=row):
                self.assertNotEqual(
                    base_hash,
                    self._preview([row]).summary.snapshot_hash,
                )

    def test_relevant_existing_metadata_change_changes_snapshot(self):
        model = self._add_model(
            "MODEL",
            vendor="Vendor",
            purchase_sku="SKU",
            cartridge_type=CartridgeType.toner,
        )
        content = _workbook_bytes(
            [("Vendor", "MODEL", "SKU", "toner", 0, None)]
        )
        first_hash = preview_cartridge_catalog_import(
            self.db,
            content,
        ).summary.snapshot_hash

        model.vendor = "Other"
        model.purchase_sku = "SKU-2"
        model.cartridge_type = CartridgeType.other
        self.db.commit()
        second_hash = preview_cartridge_catalog_import(
            self.db,
            content,
        ).summary.snapshot_hash

        self.assertNotEqual(first_hash, second_hash)

    def test_stock_columns_are_not_required_for_catalog_preview(self):
        preview = self._preview([("A", "NEW", None, "toner", 0, None)])

        self.assertEqual(preview.rows[0].status, "create")

    def test_notes_are_sanitized(self):
        preview = self._preview(
            [("A", "NEW", None, "toner", 0, "one\r\ntwo\tthree")]
        )

        self.assertEqual(preview.rows[0].notes, "one two three")

    def test_excessive_notes_are_blocking_and_not_echoed(self):
        preview = self._preview(
            [
                (
                    "A",
                    "NEW",
                    None,
                    "toner",
                    0,
                    "x" * (MAX_CATALOG_NOTES_LENGTH + 1),
                )
            ]
        )

        self.assertEqual(preview.rows[0].status, "error")
        self.assertIsNone(preview.rows[0].notes)
        self.assertTrue(any(COMMENT_HEADER in error for error in preview.rows[0].errors))

    def test_length_limits_are_blocking(self):
        preview = self._preview(
            [("V" * 101, "M" * 256, "S" * 101, "toner", 0, None)]
        )

        self.assertEqual(preview.rows[0].status, "error")
        self.assertEqual(len(preview.rows[0].errors), 3)

    def test_corrupt_xlsx_is_clear_catalog_error(self):
        with self.assertRaisesRegex(CatalogExcelError, "Не удалось прочитать XLSX-файл"):
            preview_cartridge_catalog_import(self.db, b"not-a-zip")

    def test_header_search_is_bounded_to_first_100_rows(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = INVENTORY_SHEET_NAME
        for row_number in range(1, 101):
            worksheet.cell(row=row_number, column=1, value=f"info-{row_number}")
        worksheet.append(CATALOG_HEADERS)
        worksheet.append(("A", "NEW", None, "toner", 0, None))
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()

        with self.assertRaisesRegex(CatalogExcelError, "Не найдена строка заголовков"):
            preview_cartridge_catalog_import(self.db, output.getvalue())

    def test_more_than_2000_non_empty_rows_is_error(self):
        rows = [
            ("A", f"MODEL-{index}", None, "toner", 0, None)
            for index in range(2001)
        ]

        with self.assertRaisesRegex(CatalogExcelError, "больше 2000 строк"):
            self._preview(rows)

    def test_catalog_preview_reads_xlsx_without_dimension_metadata(self):
        content = _without_dimension(
            _workbook_bytes([("A", "NEW", None, "toner", 0, None)])
        )

        preview = preview_cartridge_catalog_import(self.db, content)

        self.assertEqual(preview.rows[0].status, "create")

    def test_inventory_preview_reads_xlsx_without_dimension_metadata(self):
        self._add_model("MODEL")
        content = _without_dimension(
            _workbook_bytes(
                [("MODEL", 0, 0)],
                headers=(MODEL_NAME_HEADER, ACTUAL_NEW_HEADER, ACTUAL_REFILLED_HEADER),
            )
        )

        preview = preview_inventory_import(self.db, content)

        self.assertEqual(preview.rows[0].status, "unchanged")

    def test_preview_and_apply_routes_are_admin_only(self):
        routes = {
            route.path: route
            for route in cartridge_catalog_router.routes
            if isinstance(route, APIRoute)
        }

        for path in (
            "/api/cartridge-catalog/import/preview",
            "/api/cartridge-catalog/import/apply",
        ):
            dependency_calls = {
                dependency.call for dependency in routes[path].dependant.dependencies
            }
            self.assertIn(require_admin, dependency_calls)

    def test_apply_acquires_catalog_lock_before_preview(self):
        content = _workbook_bytes([("A", "NEW", None, "toner", 0, None)])
        preview = preview_cartridge_catalog_import(self.db, content)
        events = []

        with (
            patch(
                "app.services.cartridge_catalog_excel."
                "acquire_cartridge_catalog_mutation_lock",
                side_effect=lambda _db: events.append("lock"),
            ),
            patch(
                "app.services.cartridge_catalog_excel."
                "preview_cartridge_catalog_import",
                side_effect=lambda _db, _content: events.append("preview") or preview,
            ),
        ):
            apply_cartridge_catalog_import(
                self.db,
                content,
                preview.summary.snapshot_hash,
            )

        self.assertEqual(events, ["lock", "preview"])

    def test_regular_create_locks_before_duplicate_validation(self):
        events = []
        payload = CartridgeModelCreate(
            vendor="Vendor",
            model_name="MODEL",
            purchase_sku="SKU",
            cartridge_type=CartridgeType.toner,
        )

        with (
            patch(
                "app.api.catalog.acquire_cartridge_catalog_mutation_lock",
                side_effect=lambda _db: events.append("lock"),
            ),
            patch(
                "app.api.catalog._validate_cartridge_model_unique",
                side_effect=lambda *_args: events.append("validate"),
            ),
        ):
            post_cartridge_model(payload, self.db)

        self.assertEqual(events, ["lock", "validate"])

    def test_regular_identifier_patch_locks_before_duplicate_validation(self):
        model = self._add_model("MODEL")
        events = []

        with (
            patch(
                "app.api.catalog.acquire_cartridge_catalog_mutation_lock",
                side_effect=lambda _db: events.append("lock"),
            ),
            patch(
                "app.api.catalog._validate_cartridge_model_unique",
                side_effect=lambda *_args: events.append("validate"),
            ),
        ):
            patch_cartridge_model(
                model.id,
                CartridgeModelUpdate(model_name="UPDATED"),
                self.db,
            )

        self.assertEqual(events, ["lock", "validate"])


if __name__ == "__main__":
    unittest.main()
