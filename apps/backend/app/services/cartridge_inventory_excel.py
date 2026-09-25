from __future__ import annotations

import hashlib
import hmac
import json
import re
from collections import Counter
from io import BytesIO
from itertools import islice
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import CartridgeInventoryTransaction, CartridgeModel
from app.models.enums import CartridgeCondition, CartridgeTransactionType
from app.schemas.inventory import (
    CartridgeInventoryImportAppliedTransaction,
    CartridgeInventoryImportApplyResponse,
    CartridgeInventoryImportPreviewResponse,
    CartridgeInventoryImportPreviewRow,
    CartridgeInventoryImportPreviewSummary,
)
from app.services.cartridge_inventory import (
    acquire_cartridge_stock_mutation_lock,
    get_stock_summary,
)


INVENTORY_SHEET_NAME = "Инвентаризация"
MAX_IMPORT_ROWS = 2000
MAX_IMPORT_FILE_SIZE = 5 * 1024 * 1024
MAX_HISTORY_FILENAME_LENGTH = 180
MAX_HISTORY_ROW_COMMENT_LENGTH = 2000

MODEL_ID_HEADER = "ID модели"
VENDOR_HEADER = "Производитель"
MODEL_NAME_HEADER = "Модель картриджа *"
SKU_HEADER = "Артикул / SKU"
TYPE_HEADER = "Тип *"
CURRENT_NEW_HEADER = "Сейчас — новый"
CURRENT_REFILLED_HEADER = "Сейчас — заправленный"
ACTUAL_NEW_HEADER = "Новый — факт *"
ACTUAL_REFILLED_HEADER = "Заправленный — факт *"
INSTALLED_HEADER = "Установлено — справочно"
MIN_STOCK_HEADER = "Мин. остаток"
COMMENT_HEADER = "Комментарий"

EXPORT_HEADERS = (
    MODEL_ID_HEADER,
    VENDOR_HEADER,
    MODEL_NAME_HEADER,
    SKU_HEADER,
    TYPE_HEADER,
    CURRENT_NEW_HEADER,
    CURRENT_REFILLED_HEADER,
    ACTUAL_NEW_HEADER,
    ACTUAL_REFILLED_HEADER,
    INSTALLED_HEADER,
    MIN_STOCK_HEADER,
    COMMENT_HEADER,
)
REQUIRED_IMPORT_HEADERS = (MODEL_NAME_HEADER, ACTUAL_NEW_HEADER, ACTUAL_REFILLED_HEADER)


class InventoryExcelError(ValueError):
    pass


class InventoryApplyValidationError(InventoryExcelError):
    pass


class InventorySnapshotConflict(InventoryExcelError):
    pass


def normalize_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _normalize_header(value: object) -> str:
    normalized = normalize_text(value).replace("–", "—")
    return normalized.replace(" - ", " — ")


def build_inventory_export(db: Session) -> bytes:
    stock_by_model_id = {
        row.cartridge_model_id: row for row in get_stock_summary(db)
    }
    cartridge_models = db.scalars(
        select(CartridgeModel)
        .where(CartridgeModel.is_active.is_(True))
        .order_by(CartridgeModel.model_name)
    ).all()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = INVENTORY_SHEET_NAME

    worksheet["B1"] = "Инвентаризация склада картриджей"
    worksheet["B1"].font = Font(size=14, bold=True)
    worksheet["B2"] = (
        "Заполните колонки «Новый — факт *» и «Заправленный — факт *». "
        "Пустая ячейка не считается нулем."
    )
    worksheet["B3"] = (
        "Установленные в принтерах картриджи указаны справочно и не входят "
        "в физический остаток склада."
    )

    header_row = 5
    for column, header in enumerate(EXPORT_HEADERS, start=1):
        cell = worksheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="355C7D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_number, cartridge_model in enumerate(cartridge_models, start=header_row + 1):
        stock = stock_by_model_id[cartridge_model.id]
        values = (
            cartridge_model.id,
            cartridge_model.vendor,
            cartridge_model.model_name,
            cartridge_model.purchase_sku,
            cartridge_model.cartridge_type.value,
            stock.stock_new,
            stock.stock_refilled,
            None,
            None,
            stock.installed_total,
            cartridge_model.min_stock_level,
            cartridge_model.notes,
        )
        for column, value in enumerate(values, start=1):
            worksheet.cell(row=row_number, column=column, value=value)

    worksheet.column_dimensions["A"].hidden = True
    widths = {
        "A": 12,
        "B": 18,
        "C": 28,
        "D": 18,
        "E": 14,
        "F": 16,
        "G": 20,
        "H": 18,
        "I": 22,
        "J": 22,
        "K": 16,
        "L": 34,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "B6"
    worksheet.auto_filter.ref = f"A{header_row}:L{max(header_row, worksheet.max_row)}"
    worksheet.row_dimensions[header_row].height = 34

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def preview_inventory_import(
    db: Session,
    content: bytes,
) -> CartridgeInventoryImportPreviewResponse:
    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(content),
            data_only=True,
            read_only=True,
            keep_links=False,
        )
        if INVENTORY_SHEET_NAME not in workbook.sheetnames:
            raise InventoryExcelError(
                f"В файле отсутствует лист «{INVENTORY_SHEET_NAME}»."
            )
        worksheet = workbook[INVENTORY_SHEET_NAME]
        header_row, columns = _find_header_row(worksheet)
        return _build_preview(db, worksheet, header_row, columns)
    except InventoryExcelError:
        raise
    except Exception as exc:
        raise InventoryExcelError("Не удалось прочитать XLSX-файл.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def apply_inventory_import(
    db: Session,
    content: bytes,
    expected_snapshot_hash: str,
    filename: str,
    created_by_user_id: int,
) -> CartridgeInventoryImportApplyResponse:
    try:
        acquire_cartridge_stock_mutation_lock(db)
        preview = preview_inventory_import(db, content)

        error_rows = [row.row_number for row in preview.rows if row.status == "error"]
        if error_rows:
            row_numbers = ", ".join(str(row_number) for row_number in error_rows)
            raise InventoryApplyValidationError(
                f"Apply запрещен: файл содержит ошибки в строках {row_numbers}."
            )

        recalculated_hash = preview.summary.snapshot_hash
        if (
            re.fullmatch(r"[0-9a-f]{64}", expected_snapshot_hash) is None
            or not hmac.compare_digest(expected_snapshot_hash, recalculated_hash)
        ):
            raise InventorySnapshotConflict(
                "Остатки изменились после предварительного просмотра. "
                "Выполните Preview повторно."
            )

        pending: list[
            tuple[
                CartridgeInventoryTransaction,
                CartridgeInventoryImportPreviewRow,
                CartridgeCondition,
                str,
            ]
        ] = []
        for row in preview.rows:
            if row.status != "change":
                continue
            if row.cartridge_model_id is None or row.model_name is None:
                raise InventoryApplyValidationError(
                    f"Apply запрещен: строка {row.row_number} не сопоставлена с моделью."
                )

            _append_correction(
                db,
                pending,
                row,
                CartridgeCondition.new,
                row.delta_new,
                row.current_new,
                filename,
                created_by_user_id,
            )
            _append_correction(
                db,
                pending,
                row,
                CartridgeCondition.refilled,
                row.delta_refilled,
                row.current_refilled,
                filename,
                created_by_user_id,
            )

        db.flush()
        transactions = [
            CartridgeInventoryImportAppliedTransaction(
                transaction_id=transaction.id,
                cartridge_model_id=transaction.cartridge_model_id,
                model_name=row.model_name or "",
                condition=condition,
                direction=direction,
                quantity=transaction.quantity,
            )
            for transaction, row, condition, direction in pending
        ]
        response = CartridgeInventoryImportApplyResponse(
            snapshot_hash=recalculated_hash,
            models_processed=preview.summary.matched_rows,
            changed_models=preview.summary.changed_rows,
            transactions_created=len(transactions),
            correction_plus_total=sum(
                transaction.quantity
                for transaction, _row, _condition, direction in pending
                if direction == "plus"
            ),
            correction_minus_total=sum(
                transaction.quantity
                for transaction, _row, _condition, direction in pending
                if direction == "minus"
            ),
            transactions=transactions,
        )
        db.commit()
        return response
    except Exception:
        db.rollback()
        raise


def _append_correction(
    db: Session,
    pending: list[
        tuple[
            CartridgeInventoryTransaction,
            CartridgeInventoryImportPreviewRow,
            CartridgeCondition,
            str,
        ]
    ],
    row: CartridgeInventoryImportPreviewRow,
    condition: CartridgeCondition,
    delta: int | None,
    current: int | None,
    filename: str,
    created_by_user_id: int,
) -> None:
    if delta is None or current is None or row.cartridge_model_id is None:
        raise InventoryApplyValidationError(
            f"Apply запрещен: строка {row.row_number} содержит неполные данные."
        )
    if delta == 0:
        return

    direction = "plus" if delta > 0 else "minus"
    quantity = abs(delta)
    if direction == "minus" and quantity > current:
        raise InventorySnapshotConflict(
            f"Невозможно уменьшить остаток в строке {row.row_number}: "
            "текущий остаток недостаточен. Выполните Preview повторно."
        )

    transaction = CartridgeInventoryTransaction(
        cartridge_model_id=row.cartridge_model_id,
        transaction_type=(
            CartridgeTransactionType.correction_plus
            if direction == "plus"
            else CartridgeTransactionType.correction_minus
        ),
        quantity=quantity,
        item_condition=condition,
        reason="Инвентаризация из Excel",
        comment=_build_import_comment(filename, row.comment),
        created_by_user_id=created_by_user_id,
    )
    db.add(transaction)
    pending.append((transaction, row, condition, direction))


def _build_import_comment(filename: str, row_comment: str | None) -> str:
    basename = re.split(r"[\\/]", filename)[-1]
    safe_filename = (
        _sanitize_history_text(basename, MAX_HISTORY_FILENAME_LENGTH)
        or "inventory.xlsx"
    )
    parts = [f"Файл: {safe_filename}"]
    if row_comment:
        safe_comment = _sanitize_history_text(
            row_comment,
            MAX_HISTORY_ROW_COMMENT_LENGTH,
        )
        if safe_comment:
            parts.append(f"Комментарий Excel: {safe_comment}")
    return "; ".join(parts)


def _sanitize_history_text(value: str, max_length: int) -> str:
    cleaned = re.sub(r"[\x00-\x1f\x7f]+", " ", value)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if len(cleaned) <= max_length:
        return cleaned
    return cleaned[: max_length - 3].rstrip() + "..."


def _find_header_row(
    worksheet: Any,
    required_headers: tuple[str, ...] = REQUIRED_IMPORT_HEADERS,
) -> tuple[int, dict[str, int]]:
    required = {_normalize_header(header) for header in required_headers}
    for row_number, row in enumerate(
        islice(worksheet.iter_rows(min_row=1, values_only=True), 100),
        start=1,
    ):
        columns: dict[str, int] = {}
        for column, value in enumerate(row, start=1):
            normalized = _normalize_header(value)
            if normalized:
                columns.setdefault(normalized, column)
        if required.issubset(columns):
            return row_number, columns

    missing = ", ".join(required_headers)
    raise InventoryExcelError(
        f"Не найдена строка заголовков. Обязательные колонки: {missing}."
    )


def _build_preview(
    db: Session,
    worksheet: Any,
    header_row: int,
    columns: dict[str, int],
) -> CartridgeInventoryImportPreviewResponse:
    models = list(db.scalars(select(CartridgeModel).order_by(CartridgeModel.id)).all())
    models_by_id = {model.id: model for model in models}
    models_by_name: dict[str, list[CartridgeModel]] = {}
    for model in models:
        models_by_name.setdefault(normalize_text(model.model_name), []).append(model)

    stock_by_model_id = {
        row.cartridge_model_id: row for row in get_stock_summary(db)
    }
    rows: list[CartridgeInventoryImportPreviewRow] = []
    duplicate_keys: list[str | None] = []

    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if all(_is_blank(value) for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise InventoryExcelError(
                f"Файл содержит больше {MAX_IMPORT_ROWS} строк данных."
            )

        errors: list[str] = []
        warnings: list[str] = []
        raw_model_id = _value(values, columns, MODEL_ID_HEADER)
        raw_vendor = _optional_text(_value(values, columns, VENDOR_HEADER))
        raw_model_name = _optional_text(_value(values, columns, MODEL_NAME_HEADER))
        raw_sku = _optional_text(_value(values, columns, SKU_HEADER))
        comment = _optional_text(_value(values, columns, COMMENT_HEADER))

        if raw_model_name is None:
            errors.append(f"{MODEL_NAME_HEADER}: значение обязательно.")

        model, model_id_was_present = _match_model(
            raw_model_id,
            raw_model_name,
            raw_vendor,
            raw_sku,
            models_by_id,
            models_by_name,
            errors,
        )
        actual_new = _parse_quantity(
            _value(values, columns, ACTUAL_NEW_HEADER), ACTUAL_NEW_HEADER, errors
        )
        actual_refilled = _parse_quantity(
            _value(values, columns, ACTUAL_REFILLED_HEADER),
            ACTUAL_REFILLED_HEADER,
            errors,
        )

        stock = stock_by_model_id.get(model.id) if model else None
        current_new = stock.stock_new if stock else None
        current_refilled = stock.stock_refilled if stock else None
        installed_total = stock.installed_total if stock else None
        delta_new = (
            actual_new - current_new
            if actual_new is not None and current_new is not None
            else None
        )
        delta_refilled = (
            actual_refilled - current_refilled
            if actual_refilled is not None and current_refilled is not None
            else None
        )
        if model and not model.is_active:
            warnings.append("Модель картриджа неактивна.")
        if model and model_id_was_present:
            mismatched_identifiers = []
            if raw_model_name and normalize_text(raw_model_name) != normalize_text(
                model.model_name
            ):
                mismatched_identifiers.append(MODEL_NAME_HEADER)
            if raw_vendor and normalize_text(raw_vendor) != normalize_text(model.vendor):
                mismatched_identifiers.append(VENDOR_HEADER)
            if raw_sku and normalize_text(raw_sku) != normalize_text(model.purchase_sku):
                mismatched_identifiers.append(SKU_HEADER)
            if mismatched_identifiers:
                errors.append(
                    "MODEL ID MISMATCH: " + ", ".join(mismatched_identifiers)
                )

        rows.append(
            CartridgeInventoryImportPreviewRow(
                row_number=row_number,
                status="error",
                cartridge_model_id=model.id if model else None,
                vendor=model.vendor if model else raw_vendor,
                model_name=model.model_name if model else raw_model_name,
                purchase_sku=model.purchase_sku if model else raw_sku,
                current_new=current_new,
                actual_new=actual_new,
                delta_new=delta_new,
                current_refilled=current_refilled,
                actual_refilled=actual_refilled,
                delta_refilled=delta_refilled,
                installed_total=installed_total,
                comment=comment,
                errors=errors,
                warnings=warnings,
            )
        )
        duplicate_keys.append(
            f"id:{model.id}"
            if model
            else f"name:{normalize_text(raw_model_name)}"
            if raw_model_name
            else None
        )

    duplicate_counts = Counter(key for key in duplicate_keys if key is not None)
    for row, duplicate_key in zip(rows, duplicate_keys, strict=True):
        if duplicate_key is not None and duplicate_counts[duplicate_key] > 1:
            row.errors.append("DUPLICATE MODEL")

        if row.errors:
            row.status = "error"
        elif row.delta_new == 0 and row.delta_refilled == 0:
            row.status = "unchanged"
        else:
            row.status = "change"

    snapshot_rows = sorted(
        (
            {
                "cartridge_model_id": row.cartridge_model_id,
                "actual_new": row.actual_new,
                "actual_refilled": row.actual_refilled,
                "current_new": row.current_new,
                "current_refilled": row.current_refilled,
            }
            for row in rows
            if row.status != "error"
        ),
        key=lambda item: int(item["cartridge_model_id"] or 0),
    )
    snapshot_hash = hashlib.sha256(
        json.dumps(
            snapshot_rows,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()

    summary = CartridgeInventoryImportPreviewSummary(
        total_rows=len(rows),
        matched_rows=sum(row.cartridge_model_id is not None for row in rows),
        changed_rows=sum(row.status == "change" for row in rows),
        unchanged_rows=sum(row.status == "unchanged" for row in rows),
        error_rows=sum(row.status == "error" for row in rows),
        snapshot_hash=snapshot_hash,
    )
    return CartridgeInventoryImportPreviewResponse(summary=summary, rows=rows)


def _match_model(
    raw_model_id: object,
    raw_model_name: str | None,
    raw_vendor: str | None,
    raw_sku: str | None,
    models_by_id: dict[int, CartridgeModel],
    models_by_name: dict[str, list[CartridgeModel]],
    errors: list[str],
) -> tuple[CartridgeModel | None, bool]:
    if not _is_blank(raw_model_id):
        model_id = _parse_model_id(raw_model_id)
        if model_id is None:
            errors.append("INVALID MODEL ID")
            return None, True
        model = models_by_id.get(model_id)
        if model is None:
            errors.append("UNKNOWN MODEL")
        return model, True

    candidates = models_by_name.get(normalize_text(raw_model_name), [])
    if not candidates:
        errors.append("UNKNOWN MODEL")
        return None, False
    if len(candidates) == 1:
        return candidates[0], False

    narrowed = candidates
    if raw_vendor:
        narrowed = [
            model for model in narrowed if normalize_text(model.vendor) == normalize_text(raw_vendor)
        ]
    if raw_sku:
        narrowed = [
            model
            for model in narrowed
            if normalize_text(model.purchase_sku) == normalize_text(raw_sku)
        ]
    if len(narrowed) == 1:
        return narrowed[0], False
    errors.append("UNKNOWN MODEL" if not narrowed else "AMBIGUOUS MODEL")
    return None, False


def _parse_model_id(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float) and value.is_integer() and value > 0:
        return int(value)
    return None


def _parse_quantity(value: object, header: str, errors: list[str]) -> int | None:
    if _is_blank(value):
        errors.append(f"{header}: значение обязательно; пусто не равно нулю.")
        return None
    if isinstance(value, bool):
        errors.append(f"{header}: требуется целое число не меньше 0.")
        return None
    if isinstance(value, int):
        quantity = value
    elif isinstance(value, float) and value.is_integer():
        quantity = int(value)
    else:
        errors.append(f"{header}: требуется целое число не меньше 0.")
        return None
    if quantity < 0:
        errors.append(f"{header}: требуется целое число не меньше 0.")
        return None
    return quantity


def _value(values: tuple[Any, ...], columns: dict[str, int], header: str) -> object:
    column = columns.get(_normalize_header(header))
    if column is None or column > len(values):
        return None
    return values[column - 1]


def _optional_text(value: object) -> str | None:
    if _is_blank(value):
        return None
    return str(value).strip()


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
