from __future__ import annotations

import hashlib
import hmac
import json
import re
from collections import Counter
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import CartridgeModel
from app.models.enums import CartridgeType
from app.schemas.cartridge_catalog import (
    CartridgeCatalogImportApplyResponse,
    CartridgeCatalogImportCreatedModel,
    CartridgeCatalogImportPreviewResponse,
    CartridgeCatalogImportPreviewRow,
    CartridgeCatalogImportPreviewSummary,
)
from app.services.cartridge_catalog import acquire_cartridge_catalog_mutation_lock
from app.services.cartridge_inventory_excel import (
    COMMENT_HEADER,
    INVENTORY_SHEET_NAME,
    MAX_IMPORT_ROWS,
    MIN_STOCK_HEADER,
    MODEL_NAME_HEADER,
    SKU_HEADER,
    TYPE_HEADER,
    VENDOR_HEADER,
    InventoryExcelError,
    _find_header_row,
    _is_blank,
    _value,
    normalize_text,
)


CATALOG_REQUIRED_HEADERS = (MODEL_NAME_HEADER, TYPE_HEADER)
MAX_CATALOG_NOTES_LENGTH = 2000
TYPE_MAPPING = {
    "тонер": CartridgeType.toner,
    "toner": CartridgeType.toner,
    "чернила": CartridgeType.ink,
    "ink": CartridgeType.ink,
    "другое": CartridgeType.other,
    "other": CartridgeType.other,
}


class CatalogExcelError(ValueError):
    pass


class CatalogApplyValidationError(CatalogExcelError):
    pass


class CatalogSnapshotConflict(CatalogExcelError):
    pass


def preview_cartridge_catalog_import(
    db: Session,
    content: bytes,
) -> CartridgeCatalogImportPreviewResponse:
    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(content),
            data_only=True,
            read_only=True,
            keep_links=False,
        )
        if INVENTORY_SHEET_NAME not in workbook.sheetnames:
            raise CatalogExcelError(
                f"В файле отсутствует лист «{INVENTORY_SHEET_NAME}»."
            )
        worksheet = workbook[INVENTORY_SHEET_NAME]
        try:
            header_row, columns = _find_header_row(
                worksheet,
                CATALOG_REQUIRED_HEADERS,
            )
        except InventoryExcelError as exc:
            raise CatalogExcelError(str(exc)) from exc
        return _build_catalog_preview(db, worksheet, header_row, columns)
    except CatalogExcelError:
        raise
    except Exception as exc:
        raise CatalogExcelError("Не удалось прочитать XLSX-файл.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def apply_cartridge_catalog_import(
    db: Session,
    content: bytes,
    expected_snapshot_hash: str,
) -> CartridgeCatalogImportApplyResponse:
    try:
        acquire_cartridge_catalog_mutation_lock(db)
        preview = preview_cartridge_catalog_import(db, content)

        error_rows = [row.row_number for row in preview.rows if row.status == "error"]
        if error_rows:
            row_numbers = ", ".join(str(row_number) for row_number in error_rows)
            raise CatalogApplyValidationError(
                f"Apply запрещен: файл содержит ошибки в строках {row_numbers}."
            )

        recalculated_hash = preview.summary.snapshot_hash
        if (
            re.fullmatch(r"[0-9a-f]{64}", expected_snapshot_hash) is None
            or not hmac.compare_digest(expected_snapshot_hash, recalculated_hash)
        ):
            raise CatalogSnapshotConflict(
                "Каталог изменился после предварительного просмотра. "
                "Выполните Preview повторно."
            )

        pending: list[CartridgeModel] = []
        for row in preview.rows:
            if row.status != "create":
                continue
            if row.model_name is None or row.cartridge_type is None:
                raise CatalogApplyValidationError(
                    f"Apply запрещен: строка {row.row_number} содержит неполные данные."
                )
            pending.append(
                CartridgeModel(
                    vendor=row.vendor,
                    model_name=row.model_name,
                    purchase_sku=row.purchase_sku,
                    cartridge_type=row.cartridge_type,
                    min_stock_level=(
                        row.min_stock_level
                        if row.min_stock_level is not None
                        else 0
                    ),
                    notes=row.notes,
                    is_active=True,
                )
            )

        db.add_all(pending)
        db.flush()
        created = [
            CartridgeCatalogImportCreatedModel(
                id=model.id,
                model_name=model.model_name,
                vendor=model.vendor,
                purchase_sku=model.purchase_sku,
                cartridge_type=model.cartridge_type,
            )
            for model in pending
        ]
        response = CartridgeCatalogImportApplyResponse(
            snapshot_hash=recalculated_hash,
            models_processed=preview.summary.total_rows,
            models_created=len(created),
            existing_models=preview.summary.existing_rows,
            created=created,
        )
        db.commit()
        return response
    except Exception:
        db.rollback()
        raise


def _build_catalog_preview(
    db: Session,
    worksheet: Any,
    header_row: int,
    columns: dict[str, int],
) -> CartridgeCatalogImportPreviewResponse:
    models = list(db.scalars(select(CartridgeModel).order_by(CartridgeModel.id)).all())
    models_by_name: dict[str, list[CartridgeModel]] = {}
    models_by_sku: dict[str, list[CartridgeModel]] = {}
    for model in models:
        models_by_name.setdefault(normalize_text(model.model_name), []).append(model)
        normalized_sku = normalize_text(model.purchase_sku)
        if normalized_sku:
            models_by_sku.setdefault(normalized_sku, []).append(model)

    rows: list[CartridgeCatalogImportPreviewRow] = []
    normalized_names: list[str | None] = []
    normalized_skus: list[str | None] = []

    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if all(_is_blank(value) for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise CatalogExcelError(
                f"Файл содержит больше {MAX_IMPORT_ROWS} строк данных."
            )

        errors: list[str] = []
        warnings: list[str] = []
        model_name = _parse_text(
            _value(values, columns, MODEL_NAME_HEADER),
            MODEL_NAME_HEADER,
            errors,
            max_length=255,
            required=True,
        )
        vendor = _parse_text(
            _value(values, columns, VENDOR_HEADER),
            VENDOR_HEADER,
            errors,
            max_length=100,
        )
        purchase_sku = _parse_text(
            _value(values, columns, SKU_HEADER),
            SKU_HEADER,
            errors,
            max_length=100,
        )
        cartridge_type = _parse_cartridge_type(
            _value(values, columns, TYPE_HEADER),
            errors,
        )
        min_stock_level = _parse_min_stock(
            _value(values, columns, MIN_STOCK_HEADER),
            errors,
        )
        notes = _parse_notes(
            _value(values, columns, COMMENT_HEADER),
            errors,
        )

        normalized_name = normalize_text(model_name)
        normalized_sku = normalize_text(purchase_sku)
        candidates = models_by_name.get(normalized_name, []) if normalized_name else []
        existing_model = candidates[0] if len(candidates) == 1 else None
        if len(candidates) > 1:
            errors.append("AMBIGUOUS EXISTING MODEL")

        if existing_model is not None:
            _validate_existing_identifiers(
                existing_model,
                vendor,
                purchase_sku,
                cartridge_type,
                errors,
                warnings,
            )
            if normalized_sku:
                collision_ids = {
                    model.id for model in models_by_sku.get(normalized_sku, [])
                }
                if collision_ids - {existing_model.id}:
                    errors.append("SKU уже используется другой моделью картриджа.")
        elif normalized_sku and models_by_sku.get(normalized_sku):
            errors.append("SKU уже используется другой моделью картриджа.")

        rows.append(
            CartridgeCatalogImportPreviewRow(
                row_number=row_number,
                status="error",
                vendor=vendor,
                model_name=model_name,
                purchase_sku=purchase_sku,
                cartridge_type=cartridge_type,
                min_stock_level=min_stock_level,
                notes=notes,
                existing_model_id=existing_model.id if existing_model else None,
                errors=errors,
                warnings=warnings,
            )
        )
        normalized_names.append(normalized_name or None)
        normalized_skus.append(normalized_sku or None)

    duplicate_names = Counter(name for name in normalized_names if name is not None)
    new_skus = Counter(
        sku
        for row, sku in zip(rows, normalized_skus, strict=True)
        if row.existing_model_id is None and sku is not None
    )
    for row, normalized_name, normalized_sku in zip(
        rows,
        normalized_names,
        normalized_skus,
        strict=True,
    ):
        if normalized_name is not None and duplicate_names[normalized_name] > 1:
            row.errors.append("DUPLICATE MODEL")
        if (
            row.existing_model_id is None
            and normalized_sku is not None
            and new_skus[normalized_sku] > 1
        ):
            row.errors.append("DUPLICATE SKU")

        if row.errors:
            row.status = "error"
        elif row.existing_model_id is not None:
            row.status = "existing"
        else:
            row.status = "create"

    snapshot_hash = _build_snapshot_hash(rows, models)
    summary = CartridgeCatalogImportPreviewSummary(
        total_rows=len(rows),
        existing_rows=sum(row.status == "existing" for row in rows),
        create_rows=sum(row.status == "create" for row in rows),
        error_rows=sum(row.status == "error" for row in rows),
        snapshot_hash=snapshot_hash,
    )
    return CartridgeCatalogImportPreviewResponse(summary=summary, rows=rows)


def _parse_text(
    value: object,
    header: str,
    errors: list[str],
    *,
    max_length: int,
    required: bool = False,
) -> str | None:
    if _is_blank(value):
        if required:
            errors.append(f"{header}: значение обязательно.")
        return None
    cleaned = re.sub(r"\s+", " ", str(value).strip())
    if len(cleaned) > max_length:
        errors.append(f"{header}: длина не должна превышать {max_length} символов.")
    return cleaned


def _parse_notes(value: object, errors: list[str]) -> str | None:
    if _is_blank(value):
        return None
    cleaned = re.sub(r"[\x00-\x1f\x7f]+", " ", str(value))
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if len(cleaned) > MAX_CATALOG_NOTES_LENGTH:
        errors.append(
            f"{COMMENT_HEADER}: длина не должна превышать "
            f"{MAX_CATALOG_NOTES_LENGTH} символов."
        )
        return None
    return cleaned or None


def _parse_cartridge_type(
    value: object,
    errors: list[str],
) -> CartridgeType | None:
    normalized = normalize_text(value)
    cartridge_type = TYPE_MAPPING.get(normalized)
    if cartridge_type is None:
        errors.append(f"{TYPE_HEADER}: неизвестное значение.")
    return cartridge_type


def _parse_min_stock(value: object, errors: list[str]) -> int | None:
    if _is_blank(value):
        return 0
    if isinstance(value, bool):
        errors.append(f"{MIN_STOCK_HEADER}: требуется целое число не меньше 0.")
        return None
    if isinstance(value, int):
        result = value
    elif isinstance(value, float) and value.is_integer():
        result = int(value)
    else:
        errors.append(f"{MIN_STOCK_HEADER}: требуется целое число не меньше 0.")
        return None
    if result < 0:
        errors.append(f"{MIN_STOCK_HEADER}: требуется целое число не меньше 0.")
        return None
    return result


def _validate_existing_identifiers(
    existing_model: CartridgeModel,
    vendor: str | None,
    purchase_sku: str | None,
    cartridge_type: CartridgeType | None,
    errors: list[str],
    warnings: list[str],
) -> None:
    identifiers = (
        (VENDOR_HEADER, vendor, existing_model.vendor),
        (SKU_HEADER, purchase_sku, existing_model.purchase_sku),
    )
    for header, excel_value, db_value in identifiers:
        normalized_excel = normalize_text(excel_value)
        normalized_db = normalize_text(db_value)
        if normalized_excel and normalized_db and normalized_excel != normalized_db:
            errors.append(f"CONFLICT: {header} не совпадает с существующей моделью.")
        elif normalized_excel and not normalized_db:
            warnings.append(
                f"Модель существует; {header} из Excel не будет записан."
            )
    if (
        cartridge_type is not None
        and existing_model.cartridge_type != cartridge_type
    ):
        errors.append("CONFLICT: Тип не совпадает с существующей моделью.")


def _build_snapshot_hash(
    rows: list[CartridgeCatalogImportPreviewRow],
    models: list[CartridgeModel],
) -> str:
    snapshot = {
        "rows": [
            {
                "row_number": row.row_number,
                "status": row.status,
                "vendor": normalize_text(row.vendor),
                "model_name": normalize_text(row.model_name),
                "purchase_sku": normalize_text(row.purchase_sku),
                "cartridge_type": (
                    row.cartridge_type.value if row.cartridge_type is not None else None
                ),
                "min_stock_level": row.min_stock_level,
                "notes": row.notes,
                "existing_model_id": row.existing_model_id,
                "errors": sorted(row.errors),
                "warnings": sorted(row.warnings),
            }
            for row in rows
        ],
        "catalog": [
            {
                "id": model.id,
                "model_name": normalize_text(model.model_name),
                "vendor": normalize_text(model.vendor),
                "purchase_sku": normalize_text(model.purchase_sku),
                "cartridge_type": model.cartridge_type.value,
            }
            for model in sorted(models, key=lambda item: item.id)
        ],
    }
    return hashlib.sha256(
        json.dumps(
            snapshot,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
